"""인수인계 매뉴얼(HandoverManual)을 Word(.docx) 파일로 변환."""
import io

from docx import Document
from docx.shared import Pt, RGBColor

NAVY = RGBColor(0x1F, 0x2D, 0x3D)
MUTED = RGBColor(0x5B, 0x6B, 0x7A)


def build_manual_docx(client_name, manual):
    """HandoverManual 객체를 받아 .docx 바이트를 담은 BytesIO를 반환."""
    doc = Document()

    style = doc.styles["Normal"]
    style.font.name = "맑은 고딕"
    style.font.size = Pt(10.5)

    title = doc.add_paragraph()
    run = title.add_run(f"{client_name} — 업무 인수인계 매뉴얼")
    run.bold = True
    run.font.size = Pt(20)
    run.font.color.rgb = NAVY

    meta = doc.add_paragraph()
    meta_text = f"생성일: {manual.created_at.strftime('%Y-%m-%d %H:%M')}"
    if manual.generated_by:
        meta_text += f"  ·  인계자: {manual.generated_by.username}"
    if manual.handed_to:
        meta_text += f"  ·  인수자: {manual.handed_to.username}"
    if manual.reason_label:
        meta_text += f"  ·  사유: {manual.reason_label}"
    meta_run = meta.add_run(meta_text)
    meta_run.font.size = Pt(9.5)
    meta_run.font.color.rgb = MUTED

    if manual.handover_note:
        note_p = doc.add_paragraph()
        note_run = note_p.add_run(f"인계자 메모: {manual.handover_note}")
        note_run.italic = True
        note_run.font.size = Pt(10)
        note_run.font.color.rgb = MUTED

    doc.add_paragraph()

    for section in manual.sections:
        heading = doc.add_paragraph()
        h_run = heading.add_run(section.get("heading", ""))
        h_run.bold = True
        h_run.font.size = Pt(14)
        h_run.font.color.rgb = NAVY
        heading.paragraph_format.space_before = Pt(14)
        heading.paragraph_format.space_after = Pt(6)

        body = section.get("body", "")
        for line in body.split("\n"):
            line = line.strip()
            if not line:
                continue
            p = doc.add_paragraph(style="List Bullet" if line.startswith("- ") else None)
            p.add_run(line[2:] if line.startswith("- ") else line)

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer
